"""Build the final SIH26155 PPTX and two-page architecture PDF.

Run from the repository root:
    python3 presentation/build_submission_assets.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from weasyprint import HTML

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"
OUT.mkdir(exist_ok=True)

NAVY = RGBColor(7, 18, 35)
PANEL = RGBColor(15, 34, 55)
CYAN = RGBColor(44, 211, 225)
GREEN = RGBColor(76, 221, 166)
AMBER = RGBColor(246, 190, 65)
WHITE = RGBColor(241, 247, 251)
MUTED = RGBColor(159, 178, 196)


def textbox(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
            align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=PANEL, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def base_slide(prs, number, title, kicker):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    textbox(slide, 0.65, 0.35, 8.5, 0.25, kicker.upper(), 10, CYAN, True)
    textbox(slide, 0.65, 0.68, 11.8, 0.65, title, 28, WHITE, True)
    textbox(slide, 0.65, 7.18, 10.8, 0.18,
            "SIH26155 · NTRO · NetAudit Team · github.com/arc-1103/NetAuditor",
            8, MUTED)
    textbox(slide, 12.0, 7.12, 0.55, 0.24, f"0{number}", 11, CYAN, True,
            PP_ALIGN.RIGHT)
    return slide


def add_bullets(slide, x, y, w, bullets, size=17, gap=0.7):
    for index, text in enumerate(bullets):
        rect(slide, x, y + index * gap + 0.13, 0.12, 0.12, CYAN, False)
        textbox(slide, x + 0.28, y + index * gap, w - 0.28, 0.55, text, size)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = base_slide(prs, 1, "Mixed networks. Fragmented assurance.", "The operational gap")
    textbox(slide, 0.65, 1.55, 6.4, 1.45,
            "AI-Driven Multi-Vendor\nNetwork Security Compliance Auditor",
            28, WHITE, True)
    textbox(slide, 0.65, 3.15, 5.9, 1.0,
            "Turn raw device configuration into line-level evidence, guarded remediation and an auditable report — offline.",
            19, MUTED)
    vendors = [("CISCO", CYAN), ("FORTINET", GREEN), ("JUNIPER", AMBER),
               ("PALO ALTO", RGBColor(244, 112, 120))]
    for i, (name, color) in enumerate(vendors):
        y = 1.5 + i * 1.05
        rect(slide, 8.0, y, 2.1, 0.7, PANEL)
        textbox(slide, 8.2, y + 0.19, 1.7, 0.25, name, 12, color, True, PP_ALIGN.CENTER)
        textbox(slide, 10.25, y + 0.16, 0.6, 0.3, "→", 22, MUTED, True, PP_ALIGN.CENTER)
    rect(slide, 10.95, 2.45, 1.7, 1.7, RGBColor(17, 62, 76))
    textbox(slide, 11.18, 2.78, 1.25, 0.45, "NET", 22, CYAN, True, PP_ALIGN.CENTER)
    textbox(slide, 11.18, 3.18, 1.25, 0.35, "AUDIT", 16, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, 2, "One safe, explainable pipeline", "Architecture")
    stages = [("INGEST", "Redact + hash"), ("NORMALIZE", "AI-assisted"),
              ("VALIDATE", "Strict schema"), ("DECIDE", "OPA rules"),
              ("ACT", "Human gate"), ("REPORT", "PDF/JSON/CEF")]
    for i, (name, sub) in enumerate(stages):
        x = 0.65 + i * 2.08
        rect(slide, x, 1.65, 1.72, 1.25, PANEL)
        textbox(slide, x + 0.1, 1.95, 1.52, 0.25, name, 13, CYAN, True, PP_ALIGN.CENTER)
        textbox(slide, x + 0.1, 2.32, 1.52, 0.25, sub, 10, MUTED, False, PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            textbox(slide, x + 1.72, 2.05, 0.36, 0.25, "→", 18, MUTED, True, PP_ALIGN.CENTER)
    add_bullets(slide, 0.85, 3.55, 5.7, [
        "Local AI interprets syntax; it never decides compliance.",
        "Low confidence becomes NEEDS_REVIEW — never PASS.",
        "Only SAFE proposals can reach human approval.",
    ], 16, 0.82)
    rect(slide, 7.2, 3.55, 5.25, 2.35, RGBColor(11, 49, 60))
    textbox(slide, 7.55, 3.9, 4.55, 0.35, "THE DECISION BOUNDARY", 13, GREEN, True)
    textbox(slide, 7.55, 4.45, 4.55, 0.85,
            "AI extracts  →  Schema validates\nOPA decides  →  Human approves",
            20, WHITE, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, 3, "Evidence to safe action", "Live prototype")
    steps = [
        ("01", "Detect + protect", "Cisco IOS-XE · secrets redacted · SHA-256"),
        ("02", "Explain the risk", "11 checks · exact source lines · severity"),
        ("03", "Prepare the fix", "Reviewed CLI template · preflight status"),
        ("04", "Prove the result", "Human approval · PDF · JSON · CEF"),
    ]
    for i, (num, title, detail) in enumerate(steps):
        y = 1.48 + i * 1.28
        rect(slide, 0.72, y, 1.0, 0.82, RGBColor(13, 60, 72))
        textbox(slide, 0.87, y + 0.2, 0.7, 0.3, num, 17, CYAN, True, PP_ALIGN.CENTER)
        textbox(slide, 2.0, y + 0.02, 3.1, 0.35, title, 18, WHITE, True)
        textbox(slide, 2.0, y + 0.48, 4.3, 0.3, detail, 12, MUTED)
    rect(slide, 7.0, 1.5, 5.6, 5.0, PANEL)
    textbox(slide, 7.35, 1.87, 4.9, 0.3, "CIS-NET-1.1.2", 11, CYAN, True)
    textbox(slide, 7.35, 2.3, 4.9, 0.7, "Telnet administrative\naccess is enabled", 23, WHITE, True)
    rect(slide, 7.35, 3.28, 1.25, 0.48, RGBColor(92, 35, 43))
    textbox(slide, 7.48, 3.4, 0.98, 0.2, "CRITICAL", 10, RGBColor(255, 140, 145), True, PP_ALIGN.CENTER)
    textbox(slide, 7.35, 4.05, 4.8, 0.55, "line 31: transport input telnet ssh", 13, MUTED)
    rect(slide, 7.35, 5.15, 2.1, 0.62, RGBColor(16, 75, 61))
    textbox(slide, 7.6, 5.34, 1.6, 0.22, "SAFE TO REVIEW", 11, GREEN, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, 4, "Innovation with defensible claims", "Proof and scope")
    rows = [
        ("AI / rules separation", "Schema gate + OPA-only verdicts"),
        ("Unknown-vendor learning", "Admin-reviewed mapping and RAG APIs"),
        ("Safe remediation", "Provenance + preflight + approval"),
        ("Audit-ready output", "Evidence + versioned policy context"),
        ("Sovereign deployment", "No cloud dependency at runtime"),
    ]
    for i, (left, right) in enumerate(rows):
        y = 1.4 + i * 0.82
        rect(slide, 0.7, y, 11.9, 0.64, PANEL)
        textbox(slide, 0.95, y + 0.17, 3.8, 0.22, left, 13, WHITE, True)
        textbox(slide, 4.8, y + 0.17, 6.9, 0.22, right, 13, MUTED)
        textbox(slide, 11.8, y + 0.14, 0.42, 0.25, "✓", 15, GREEN, True, PP_ALIGN.CENTER)
    rect(slide, 0.7, 5.8, 11.9, 0.68, RGBColor(57, 43, 17))
    textbox(slide, 0.95, 5.99, 11.3, 0.28,
            "TODAY: Cisco + Fortinet · single-file UI · one 11-control CIS-inspired bundle · not CIS certification",
            12, AMBER, True, PP_ALIGN.CENTER)

    slide = base_slide(prs, 5, "From prototype to controlled pilot", "Scale path and impact")
    phases = [
        ("30 DAYS", "Bulk ingestion\nJuniper + Palo Alto"),
        ("60 DAYS", "Authorized framework bundles\nLearning-queue UI"),
        ("90 DAYS", "Sanitized pilot\nPrecision + time metrics"),
    ]
    for i, (time, body) in enumerate(phases):
        x = 0.72 + i * 4.12
        rect(slide, x, 1.5, 3.72, 2.15, PANEL)
        textbox(slide, x + 0.25, 1.82, 3.22, 0.3, time, 12, CYAN, True)
        textbox(slide, x + 0.25, 2.35, 3.22, 0.85, body, 18, WHITE, True)
    textbox(slide, 0.72, 4.15, 12.0, 0.55,
            "One explainable workflow · Less manual effort · Faster evidence · Safer change preparation",
            19, GREEN, True, PP_ALIGN.CENTER)
    rect(slide, 1.45, 5.05, 10.4, 1.05, RGBColor(11, 49, 60))
    textbox(slide, 1.75, 5.28, 9.8, 0.25,
            "ASK  ·  Sanitized configurations  ·  Framework SMEs  ·  Controlled pilot",
            14, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 1.75, 5.7, 9.8, 0.22,
            "github.com/arc-1103/NetAuditor  ·  Manav Mishra  ·  manavmishra260205@gmail.com",
            10, MUTED, False, PP_ALIGN.CENTER)

    path = OUT / "NetAudit_SIH26155_Technical_Presentation.pptx"
    prs.save(path)
    return path, len(prs.slides)


def build_pdf():
    css = """
    @page { size: A4; margin: 13mm 14mm; @bottom-right { content: "NetAudit · SIH26155 · " counter(page) "/" counter(pages); font: 8pt sans-serif; color: #526272; } }
    * { box-sizing: border-box; } body { font-family: Arial, sans-serif; color: #102235; font-size: 9.2pt; line-height: 1.34; margin: 0; }
    h1 { color: #087b8c; font-size: 22pt; margin: 0 0 2mm; } h2 { color: #087b8c; font-size: 13pt; margin: 4mm 0 1.5mm; }
    p { margin: 1.5mm 0; } .meta { color: #526272; font-weight: bold; margin-bottom: 4mm; }
    .box { background: #eef7f8; border-left: 3px solid #10a6b9; padding: 3mm; margin: 2mm 0; }
    pre { background: #102235; color: #dffcff; padding: 3mm; font: 7.8pt monospace; line-height: 1.25; white-space: pre-wrap; }
    ol { padding-left: 5mm; margin: 2mm 0; } li { margin: 1.3mm 0; }
    table { width: 100%; border-collapse: collapse; font-size: 8.2pt; margin-top: 2mm; } th { background: #102235; color: white; } th, td { border: 0.3mm solid #b8c7d2; padding: 1.8mm; vertical-align: top; }
    .page-break { break-before: page; } .small { font-size: 8pt; color: #526272; }
    """
    html = f"""<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>
    <h1>NetAudit Engine — Submission Architecture</h1>
    <div class='meta'>SIH26155 · NTRO · AI-Driven Multi-Vendor Network Security Compliance Auditor</div>
    <div class='box'><b>Safety boundary:</b> local AI may interpret unfamiliar syntax, but only version-controlled OPA rules produce PASS/FAIL. Invalid or low-confidence extraction becomes <b>NEEDS_REVIEW</b>, never compliant.</div>
    <h2>System architecture</h2>
    <pre>Browser → API Gateway → Ingestion → Redis → Parsing Worker
                            ↓                  ↓
                   MinIO + PostgreSQL   SecurityBaseline JSON
                                               ↓
                                      OPA Compliance Engine
                                         ↓              ↓
                               Remediation + Gate   PDF / JSON / CEF

Optional profiles: reviewed learning/RAG · topology blast radius · Batfish · Ollama</pre>
    <p>Runtime services use an internal Docker network; only the dashboard and gateway publish host ports. Uploads are bounded, MIME checked, credential-redacted, SHA-256 fingerprinted and stored with provenance. RBAC separates admin, operator and auditor access.</p>
    <h2>Five-stage pipeline</h2>
    <ol>
      <li><b>Ingest:</b> accept configuration text; validate, redact, hash and enqueue.</li>
      <li><b>Normalize:</b> detect vendor/OS and map CLI to a strict Pydantic SecurityBaseline. Cisco IOS-XE and Fortinet FortiOS are demonstrated.</li>
      <li><b>Learn:</b> unknown blocks enter an administrator-reviewed mapping queue. Retrieved context never changes a verdict directly.</li>
      <li><b>Evaluate:</b> OPA applies an 11-control CIS-inspired prototype bundle and retains line-level evidence, severity and versions.</li>
      <li><b>Act and report:</b> reviewed Jinja2 CLI is preflighted; only SAFE proposals can be approved. Evidence exports as HTML/PDF/JSON/CEF.</li>
    </ol>
    <h2>Extensibility model</h2>
    <p>New vendors reuse the normalized schema and service pipeline; they add validated examples/mappings and reviewed remediation templates. New frameworks are versioned OPA bundles over the stable schema, isolating policy updates from vendor syntax.</p>

    <div class='page-break'></div>
    <h1>Prototype Evidence, Limits and Deployment</h1>
    <table><thead><tr><th>SIH requirement</th><th>Current proof</th><th>Next validated milestone</th></tr></thead><tbody>
      <tr><td>Unified ingestion</td><td>Dashboard upload, redaction, SHA-256 and MinIO</td><td>Bulk inventory UI and orchestration</td></tr>
      <tr><td>AI adaptation</td><td>Local structured parsing and reviewed learning APIs</td><td>Wire admin learning queue into dashboard</td></tr>
      <tr><td>Multi-framework engine</td><td>Pluggable, versioned OPA architecture</td><td>Authorized CIS/NIST/STIG/ISO content</td></tr>
      <tr><td>Actionable reporting</td><td>Evidence, severity, remediation, PDF/JSON/CEF</td><td>Validate more hardware and firmware</td></tr>
      <tr><td>Vendor-agnostic scale</td><td>Shared schema; Cisco and Fortinet fixtures</td><td>Juniper and Palo Alto validation</td></tr>
    </tbody></table>
    <h2>Deployment and trust</h2>
    <p>The default Compose profile is the deterministic presentation core. Optional <b>advanced</b> and <b>model</b> profiles add learning/topology/Batfish and local Ollama. The system does not push configuration to devices automatically: a human remains responsible for review and change control.</p>
    <h2>Verification evidence</h2>
    <div class='box'>Gateway 25 passed · Ingestion 20 passed · Compliance 98 passed (6 external integration skips) · Remediation 55 passed · Reporting 14 passed · Schema 8 passed · Scripts 6 passed · frontend type-check and production build passed.</div>
    <h2>Prototype claim boundary</h2>
    <p>This prototype demonstrates two vendors, a single-file UI and one 11-control CIS-inspired policy bundle. It does <b>not</b> claim CIS certification or complete CIS, NIST, STIG or ISO coverage. The architecture is designed for those additions after content authorization and validation against representative sanitized configurations.</p>
    <h2>Submission references</h2>
    <p><b>Repository:</b> https://github.com/arc-1103/NetAuditor<br><b>Contact:</b> Manav Mishra · manavmishra260205@gmail.com</p>
    <p class='small'>Generated from the tested submission branch. See README.md and docs/SUBMISSION_CHECKLIST.md for setup and evidence guidance.</p>
    </body></html>"""
    path = OUT / "NetAudit_SIH26155_Architecture.pdf"
    document = HTML(string=html, base_url=str(ROOT)).render()
    document.write_pdf(path)
    return path, len(document.pages)


if __name__ == "__main__":
    pptx_path, slides = build_pptx()
    pdf_path, pages = build_pdf()
    print(f"PPTX: {pptx_path} ({slides} slides)")
    print(f"PDF:  {pdf_path} ({pages} pages)")
